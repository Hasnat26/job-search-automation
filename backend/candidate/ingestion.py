from hashlib import sha256
from pathlib import Path
from typing import Any

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def validate_cv_path(path: str | Path) -> Path:
    candidate = Path(path).expanduser().resolve()
    if not candidate.exists():
        raise FileNotFoundError(candidate)
    if not candidate.is_file():
        raise ValueError(f"CV path is not a file: {candidate}")
    if candidate.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported CV format: {candidate.suffix}")
    return candidate


def sha256_file(path: str | Path) -> str:
    candidate = validate_cv_path(path)
    digest = sha256()
    with candidate.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_text(path: str | Path) -> tuple[str, dict[str, Any]]:
    candidate = validate_cv_path(path)
    suffix = candidate.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = candidate.read_text(encoding="utf-8", errors="replace")
        return text, {"page_count": None, "extractor": "plain_text"}
    if suffix == ".pdf":
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:
            raise RuntimeError("PDF support requires PyMuPDF (pip install -r requirements.txt)") from exc
        with fitz.open(candidate) as doc:
            pages = [page.get_text("text") for page in doc]
            text = "\n\n".join(pages)
            return text, {"page_count": len(pages), "extractor": "pymupdf"}
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx (pip install -r requirements.txt)") from exc
        doc = Document(candidate)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs), {"page_count": None, "extractor": "python-docx"}
    raise ValueError(f"Unsupported CV format: {suffix}")
